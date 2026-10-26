"""Verify integration of accepted provenance and CAP5 evidence without rerunning statistics."""
from __future__ import annotations
import argparse
import collections
import csv
import hashlib
import io
import json
import subprocess
import zipfile
from pathlib import Path

ROOT=Path(__file__).resolve().parents[2]
REPO=ROOT.parent
BASELINE='7ef3b25fddb32b3d7089db90974fdfccb405847f'
P=ROOT/'results/current/provenance_audit_20261023'
C=ROOT/'results/current/chemical_diversity_cap5_20261023'


def read(path):
    with path.open(newline='',encoding='utf-8') as f:return list(csv.DictReader(f))


def main():
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--baseline-ref',default=BASELINE);args=ap.parse_args()
    checks=[]
    def check(name,ok,detail=''):checks.append({'check':name,'status':'PASS' if ok else 'FAIL','detail':detail})
    def old(path):
        return subprocess.check_output(['git','show',f'{args.baseline_ref}:{path.relative_to(REPO).as_posix()}'],cwd=REPO)
    # Protect the formal result cohort, predictions, ranks, source rows, CAP5
    # accepted outputs and existing primary images against integration edits.
    protected=[]
    for rel in ['results/formal_metrics','results/jcim_stage2','results/current/provenance_audit_20261023','results/current/chemical_diversity_cap5_20261023','activity_label_source_audit_20261008']:
        protected.extend(p for p in (ROOT/rel).iterdir() if p.is_file())
    protected += [ROOT/x for x in ['results/current/PRIMARY_DIRECTIONAL_ANALYSIS_POPULATION.csv','results/current/pair_ligand_mapping.csv','results/canonical/model_fold_assignments.csv','analysis_plan_jcim/CHEMISTRY_POPULATION_LAYERS.csv','manuscript/Main_Tables.xlsx']]
    protected += sorted((ROOT/'results/jcim_stage2_figures').glob('*.png'))
    for p in protected:check('protected_bytes:'+p.relative_to(ROOT).as_posix(),p.read_bytes()==old(p))
    pop=read(ROOT/'results/current/PRIMARY_DIRECTIONAL_ANALYSIS_POPULATION.csv');eligible=[r for r in pop if r['activity_eligible']=='1' and r['class']]
    check('main_population_867',len(pop)==869 and len(eligible)==867)
    check('main_classes_unchanged',collections.Counter(r['class'] for r in eligible)=={'dual':248,'A_only':256,'B_only':252,'neither':111})
    source=read(P/'CURRENT_867_ACTIVITY_SOURCE_CROSSWALK.csv');keys=lambda rows:{(r['pair'],r['panel_id']) for r in rows}
    check('source_covers_exact_main_members',len(source)==867 and keys(source)==keys(eligible))
    check('source_both_ends',all(r['both_ends_raw_activity']=='YES' for r in source))
    main_classes={(r['pair'],r['panel_id']):r['class'] for r in eligible}
    check('source_classes_match_main',all(r['class']==main_classes[r['pair'],r['panel_id']] for r in source))
    expected={'literature_complete':855,'document_identifier_missing':12,'cross_source_or_endpoint_mix':168,'replicate_threshold_discordance':35,'value_or_label_mismatch':0}
    for flag,n in expected.items():check('source_flag:'+flag,sum(r[flag]=='YES' for r in source)==n)
    audit=read(P/'AUDIT_PASSFAIL.csv');check('source_no_fail',not any(r['result']=='FAIL' for r in audit));check('selection_unverifiable_retained',sum(r['result']=='NOT_VERIFIABLE' for r in audit)==2)
    replay=read(P/'SEVEN_PAIR_EXACT_REPLAY_RESULTS.csv');statusfields=[k for k in replay[0] if any('DOCUMENTED_NOT_REPLAYED'==r[k] for r in replay)]
    check('seven_selection_statuses',len(replay)==7 and len(statusfields)==1 and all(r[statusfields[0]]=='DOCUMENTED_NOT_REPLAYED' for r in replay))
    membership=read(C/'CAP5_MEMBERSHIP_867.csv');retained=[r for r in membership if r['retained_cap5']=='YES'];excluded=[r for r in membership if r['retained_cap5']!='YES']
    check('CAP5_membership',keys(membership)==keys(eligible) and len(retained)==861 and len(excluded)==6)
    check('CAP5_exclusions',collections.Counter((r['pair'],r['class']) for r in excluded)=={('JAK1/JAK2','A_only'):3,('PPARA/PPARD','A_only'):1,('PPARA/PPARD','B_only'):2})
    check('CAP5_neither_111',sum(r['class']=='neither' for r in retained)==111)
    check('CAP5_acyclic_four_retained',sum(r['scaffold_is_acyclic']=='YES' for r in retained)==4)
    counts=collections.Counter((r['pair'],r['class'],r['murcko_scaffold']) for r in retained if r['scaffold_is_acyclic']!='YES')
    check('CAP5_max_repeat',max(counts.values())<=5)
    check('CAP5_no_failure',all(r['result']=='PASS' for r in read(C/'CAP5_AUDIT_PASSFAIL.csv')))
    compare=read(C/'CAP5_PRIMARY_VS_ORIGINAL.csv');check('comparison_336',len(compare)==336)
    check('comparison_point_classes',collections.Counter(r['point_change_class'] for r in compare)=={'STABLE':278,'NEAR_NULL':42,'AMPLIFIED':9,'ATTENUATED':7})
    # Compare displayed M0 and chemistry claims with accepted tables.
    primary=read(ROOT/'results/formal_metrics/PRIMARY_DIRECTIONAL_METRICS.csv');capprimary=read(C/'CAP5_PRIMARY_DIRECTIONAL_METRICS.csv')
    pm={(r['pair'],r['method'],r['contrast']):r for r in primary};cm={(r['pair'],r['method'],r['contrast']):r for r in capprimary}
    for pair,value in [('JAK1/JAK2',0.6314655172),('PPARA/PPARD',0.496875),('PIK3CA/mTOR',0.580078125)]:check('CAP5_M0_summary:'+pair,abs(float(cm[pair,'M0','summary_min']['auroc'])-value)<1e-9)
    units=read(ROOT/'results/jcim_stage2/module_D_chemistry_units.csv');cu=read(C/'CAP5_MODULE_D_CHEMISTRY_UNITS.csv');by={(r['pair'],r['arm']):r for r in units}
    check('ECFP4_all_changes_under_002',len(cu)==16 and all(abs(float(r['auroc_ecfp4'])-float(by[r['pair'],r['arm']]['auroc_ecfp4']))<.02 for r in cu))
    from openpyxl import load_workbook
    oldbook=load_workbook(io.BytesIO(old(ROOT/'manuscript/Supplementary_Tables.xlsx')),read_only=True,data_only=True)
    book=load_workbook(ROOT/'manuscript/Supplementary_Tables.xlsx',read_only=True,data_only=True)
    for name in oldbook.sheetnames:check('existing_sheet_values:'+name,list(oldbook[name].values)==list(book[name].values))
    from jcim_manuscript_integrate import EXTRA_SHEETS
    for name,path in EXTRA_SHEETS:
        source_rows=read(path);fields=list(source_rows[0]);expected=[fields]+[[r.get(k,'') for k in fields] for r in source_rows]
        actual=[[str(v) if v is not None else '' for v in row[:len(fields)]] for row in book[name].values if any(v is not None and v!='' for v in row)]
        check('supporting_sheet_matches_csv:'+name,expected==actual)
    check('supplementary_sheet_count',len(book.sheetnames)==31);oldbook.close();book.close()
    with zipfile.ZipFile(ROOT/'manuscript/Supplementary_Source_Data.zip') as z:
        mf=list(csv.DictReader(io.StringIO(z.read('SHA256SUMS.csv').decode())));listed={r['relative_path'] for r in mf}
        check('package_checksums',all(hashlib.sha256(z.read(r['relative_path'])).hexdigest()==r['sha256'] for r in mf))
        check('package_current_bytes',all(z.read(r['relative_path'])==(REPO/r['relative_path']).read_bytes() for r in mf))
        check('package_all_new_evidence',all(p.relative_to(REPO).as_posix() in listed for d in [P,C] for p in d.iterdir() if p.is_file() and p.suffix in {'.csv','.json','.md','.py'}))
        check('package_no_PM48_regression_outputs',not any('/regression_pm48/' in p for p in z.namelist()))
    authority=json.loads((ROOT/'CURRENT_RESULT_AUTHORITIES.json').read_text());check('CAP5_not_promoted',authority['cap5']['replaces_primary'] is False and authority['cap5']['role']=='supporting_sensitivity_only')
    # Obsolete missing-source assertions must disappear from active text.
    for name in ['Supporting_Information_CN.md','Source_Register.csv']:
        t=(ROOT/'manuscript'/name).read_text();check('no_stale_missing_export:'+name,'Q110 未提供完整逐 assay 来源导出' not in t and '完整 Q110 逐 assay 来源导出仍缺失' not in t)
    import fitz
    for name in ['Results_Review.pdf','Supporting_Information_Review.pdf']:
        d=fitz.open(ROOT/'manuscript'/name);text=''.join(p.get_text() for p in d);check('PDF_evidence_updated:'+name,'861' in text and '855' in text);d.close()
    report={'source_commit':BASELINE,'scientific_recomputation':False,'n_checks':len(checks),'n_pass':sum(r['status']=='PASS' for r in checks),'n_fail':sum(r['status']=='FAIL' for r in checks),'source_audit_not_verifiable':2,'checks':checks}
    (ROOT/'results/current/CURRENT_VERIFICATION_REPORT.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps({k:v for k,v in report.items() if k!='checks'}))
    for r in checks:
        if r['status']=='FAIL':print(r)
    return bool(report['n_fail'])


if __name__=='__main__':raise SystemExit(main())

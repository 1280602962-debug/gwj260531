"""Build manuscript tables and the source-data package from accepted files.

Read-only with respect to scientific inputs. Writes only under manuscript/.
Does not recompute AUROC, intervals, predictions, or ranks.
"""

from __future__ import annotations

import csv
import hashlib
import shutil
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
REPO = ROOT.parent
MAN = ROOT / "manuscript"
VIEWS = MAN / "source_views"
ORDER = [
    "EGFR/HER2",
    "JAK1/JAK2",
    "JAK1/TYK2",
    "PIK3CA/mTOR",
    "AChE/BChE",
    "F2/F10",
    "PPARG/PPARA",
    "PPARA/PPARD",
]


def read_csv(path: Path) -> tuple[list[str], list[dict]]:
    with path.open(newline="") as handle:
        reader = csv.DictReader(handle)
        rows = list(reader)
        return list(reader.fieldnames or []), rows


def write_sheet(ws, fields: list[str], rows: list[dict]) -> None:
    header = Font(name="Arial", bold=True, size=9)
    ws.append(fields)
    for cell in ws[1]:
        cell.font = header
    for row in rows:
        ws.append(["" if row.get(field) is None else str(row.get(field)) for field in fields])
    for index, field in enumerate(fields, start=1):
        ws.column_dimensions[get_column_letter(index)].width = min(28, max(12, len(field) + 2))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def build_review_pdfs() -> None:
    import tempfile
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from fontTools.ttLib import TTCollection, TTFont
    from matplotlib import font_manager
    from matplotlib.backends.backend_pdf import PdfPages
    from PIL import Image

    with tempfile.TemporaryDirectory() as tmp:
        font_path = Path(tmp) / "msyh.ttf"
        TTCollection("/mnt/c/Windows/Fonts/msyh.ttc").fonts[0].save(font_path)
        face = TTFont(font_path)
        cmap = face.getBestCmap()
        advances = face["hmtx"].metrics
        upem = face["head"].unitsPerEm

        def char_width(char: str, size: float) -> float:
            glyph = cmap.get(ord(char))
            if glyph is None:
                return size * 0.5
            return advances[glyph][0] / upem * size

        def wrap(paragraph: str, size: float, max_pt: float) -> list[str]:
            if not paragraph:
                return [""]
            tokens: list[str] = []
            latin = ""
            for char in paragraph:
                if ord(char) < 128 and char != " ":
                    latin += char
                    continue
                if latin:
                    tokens.append(latin)
                    latin = ""
                tokens.append(char)
            if latin:
                tokens.append(latin)
            lines: list[str] = []
            current = ""
            width = 0.0
            for token in tokens:
                advance = sum(char_width(char, size) for char in token)
                if current and width + advance > max_pt:
                    lines.append(current.rstrip())
                    current = token.lstrip() if token != " " else ""
                    width = sum(char_width(char, size) for char in current)
                else:
                    current += token
                    width += advance
            if current.strip():
                lines.append(current.rstrip())
            return lines or [""]

        def render(path: Path, blocks: list[tuple[str, Path | None]]) -> None:
            page_w, page_h = 8.27, 11.69
            margin = 0.55
            max_pt = (page_w - 2 * margin) * 72
            pdf = PdfPages(path)
            fig = None
            y = 0.0

            def new_page():
                nonlocal fig, y
                if fig is not None:
                    pdf.savefig(fig)
                    plt.close(fig)
                fig = plt.figure(figsize=(page_w, page_h))
                y = page_h - margin

            def ensure(need: float) -> None:
                if y - need < margin:
                    new_page()

            new_page()
            prop = font_manager.FontProperties(fname=str(font_path))

            def write_text(payload: str) -> None:
                nonlocal y
                for paragraph in payload.splitlines():
                    size = 10.5
                    leading = 14.5 / 72
                    for line in wrap(paragraph, size, max_pt):
                        ensure(leading)
                        fig.text(margin / page_w, y / page_h, line, fontproperties=prop, fontsize=size, va="top")
                        y -= leading
                    y -= 2 / 72

            def image_size(image_path: Path) -> tuple[float, float]:
                image = Image.open(image_path)
                iw, ih = image.size
                draw_w = page_w - 2 * margin
                draw_h = draw_w * ih / iw
                max_h = page_h - 2 * margin
                if draw_h > max_h:
                    draw_h = max_h
                    draw_w = draw_h * iw / ih
                return draw_w, draw_h

            index = 0
            while index < len(blocks):
                kind, payload = blocks[index]
                nxt = blocks[index + 1] if index + 1 < len(blocks) else None
                if kind == "text" and nxt and nxt[0] == "image" and "\n" not in str(payload).strip():
                    title = str(payload)
                    draw_w, draw_h = image_size(nxt[1])
                    need = 16 / 72 + draw_h + 0.15
                    if y - need < margin:
                        new_page()
                    write_text(title)
                    image = Image.open(nxt[1])
                    y -= draw_h
                    ax = fig.add_axes([margin / page_w, y / page_h, draw_w / page_w, draw_h / page_h])
                    ax.imshow(image)
                    ax.axis("off")
                    y -= 0.15
                    index += 2
                    continue
                if kind == "text":
                    write_text(str(payload))
                else:
                    draw_w, draw_h = image_size(payload)
                    ensure(draw_h + 0.15)
                    image = Image.open(payload)
                    y -= draw_h
                    ax = fig.add_axes([margin / page_w, y / page_h, draw_w / page_w, draw_h / page_h])
                    ax.imshow(image)
                    ax.axis("off")
                    y -= 0.15
                index += 1
            if fig is not None:
                pdf.savefig(fig)
                plt.close(fig)
            pdf.close()

        fig_dir = ROOT / "results" / "jcim_stage2_figures"
        results_blocks: list[tuple[str, Path | None]] = [
            ("text", "结果材料审阅稿\n\n本文件合并 Results、Scheme 1、六张结果主图、两张正文表和图注，供阅读核对。它不是完整投稿稿，不含 Introduction、完整 Methods 和 Discussion。\n\n"),
            ("text", (MAN / "Results_CN.md").read_text() + "\n"),
            ("text", "Table 1 与 Table 2 的完整单元格在 Main_Tables.xlsx。下图为正式矢量图的预览嵌入，排版母版为同名 PDF，宽度 177.8 mm。\n"),
        ]
        for stem, title in [
            ("Scheme01_Study_Design", "Scheme 1"),
            ("Fig01_Primary_Directional_and_Deltas", "Figure 1"),
            ("Fig02_Negative_Class_Contrast", "Figure 2"),
            ("Fig03_Chemistry_and_Docking_Increment", "Figure 3"),
            ("Fig04_Pocket_Contrast", "Figure 4"),
            ("Fig05_Seed_and_Receptor_Sensitivity", "Figure 5"),
            ("Fig06_Directional_and_Topk_Changes", "Figure 6"),
        ]:
            results_blocks.append(("text", title))
            results_blocks.append(("image", fig_dir / f"{stem}.png"))
        results_blocks.append(("text", (MAN / "Figure_Captions_CN.md").read_text()))
        render(MAN / "Results_Review.pdf", results_blocks)
        si_blocks: list[tuple[str, Path | None]] = [
            ("text", "补充信息审阅稿\n\n长表不在本 PDF 逐行排印。Tables S1–S17 见 Supplementary_Tables.xlsx，原 CSV 见 Supplementary_Source_Data.zip。\n\n"),
            ("text", (MAN / "Supporting_Information_CN.md").read_text() + "\n"),
        ]
        for stem, title in [
            ("FigS01_Descriptors_D_vs_A", "Figure S1"),
            ("FigS02_Descriptors_D_vs_B", "Figure S2"),
            ("FigS03_Cross_Class_Similarity", "Figure S3"),
            ("FigS04_All_Method_AUROC", "Figure S4"),
            ("FigS05_Five_Seeds", "Figure S5"),
            ("FigS06_Alternative_Receptors", "Figure S6"),
            ("FigS07_Missingness", "Figure S7"),
            ("FigS08_Topk_and_Concordance", "Figure S8"),
        ]:
            si_blocks.append(("text", title))
            si_blocks.append(("image", fig_dir / f"{stem}.png"))
        si_blocks.append(("text", (MAN / "Figure_Captions_CN.md").read_text()))
        render(MAN / "Supporting_Information_Review.pdf", si_blocks)
        print("REVIEW_PDFS", MAN / "Results_Review.pdf", MAN / "Supporting_Information_Review.pdf")


def main() -> None:
    MAN.mkdir(parents=True, exist_ok=True)
    VIEWS.mkdir(parents=True, exist_ok=True)
    pop_fields, pop = read_csv(ROOT / "reruns/UNIFORM_RERUN_V4_2_20260921/13_qa/PRIMARY_DIRECTIONAL_ANALYSIS_POPULATION.csv")
    map_fields, mapping = read_csv(ROOT / "reruns/UNIFORM_RERUN_V4_2_20260921/00_protocol/pair_ligand_mapping.csv")
    eligible = [r for r in pop if r["activity_eligible"] == "1" and r["class"] in {"dual", "A_only", "B_only", "neither"}]
    main = Workbook()
    ws = main.active
    ws.title = "Table1_Label_Counts"
    fields = ["index", "pair", "target_A", "pdb_A", "target_B", "pdb_B", "dual", "A_only", "B_only", "neither"]
    ws.append(fields)
    for index, pair in enumerate(ORDER, start=1):
        rows = [r for r in eligible if r["pair"] == pair]
        pdb = sorted({(r["pdb_A"], r["pdb_B"]) for r in rows})
        if len(pdb) != 1:
            raise SystemExit(f"PDB not unique for {pair}: {pdb}")
        target_a, target_b = pair.split("/")
        counts = {name: sum(r["class"] == name for r in rows) for name in ("dual", "A_only", "B_only", "neither")}
        ws.append([index, pair, target_a, pdb[0][0], target_b, pdb[0][1], counts["dual"], counts["A_only"], counts["B_only"], counts["neither"]])
    desc_fields, desc = read_csv(ROOT / "results/jcim_stage2/module_I_descriptive.csv")
    by = {(r["pair"], r["method"], r["population_kind"]): r for r in desc}
    ws2 = main.create_sheet("Table2_Common_Topk")
    header = ["index", "pair"]
    for method in ("M1", "M1b", "M2", "M3"):
        tag = f"{method}_supporting" if method == "M1b" else method
        header.extend([f"{tag}_N_common", f"{tag}_k", f"{tag}_dual_M0", f"{tag}_dual_new"])
    ws2.append(header)
    for index, pair in enumerate(ORDER, start=1):
        values = [index, pair]
        for method in ("M1", "M1b", "M2", "M3"):
            base = by[pair, "M0", f"pairwise_common_vs_{method}"]
            new = by[pair, method, "pairwise_common_vs_M0"]
            if base["n_total"] != new["n_total"] or base["k"] != new["k"]:
                raise SystemExit(f"common-list mismatch {pair} {method}")
            values.extend([base["n_total"], base["k"], base["n_dual_top"], new["n_dual_top"]])
        ws2.append(values)
    for sheet in main.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                cell.font = Font(name="Arial", size=9, bold=(cell.row == 1))
                cell.number_format = "@" if cell.row > 1 else "General"
    main.save(MAN / "Main_Tables.xlsx")

    sheets = [
        ("S1_Mapping", ROOT / "reruns/UNIFORM_RERUN_V4_2_20260921/00_protocol/pair_ligand_mapping.csv"),
        ("S2_Population_Layers", ROOT / "analysis_plan_jcim/CHEMISTRY_POPULATION_LAYERS.csv"),
        ("S3_PRIMARY_Metrics", ROOT / "results/formal_metrics/PRIMARY_DIRECTIONAL_METRICS.csv"),
        ("S4_PRIMARY_Deltas", ROOT / "results/formal_metrics/PRIMARY_METHOD_DELTA_METRICS.csv"),
        ("S5_Negative_Class", ROOT / "results/jcim_stage2/module_B_negative_class.csv"),
        ("S6_Wrong_Pocket", ROOT / "results/jcim_stage2/module_C_wrong_pocket.csv"),
        ("S7_Chemistry_Units", ROOT / "results/jcim_stage2/module_D_chemistry_units.csv"),
        ("S8_Composition", ROOT / "results/jcim_stage2/module_D1_composition.csv"),
        ("S9_Chemistry_OOF", ROOT / "results/jcim_stage2/module_D_chemistry_oof_ligands.csv"),
        ("S10_Five_Seeds", ROOT / "results/jcim_stage2/module_E_five_seed.csv"),
        ("S11_Seed_Summaries", ROOT / "results/jcim_stage2/module_E_five_seed_summary.csv"),
        ("S12_Alt_Receptors", ROOT / "results/jcim_stage2/module_F_alt_receptor.csv"),
        ("S13_Missingness", ROOT / "results/jcim_stage2/module_H_missingness.csv"),
        ("S14_Ranking_Descriptive", ROOT / "results/jcim_stage2/module_I_descriptive.csv"),
        ("S15_Ranking_Comparison", ROOT / "results/jcim_stage2/module_I_method_compare.csv"),
        ("S16_Ligand_Ranks", ROOT / "results/jcim_stage2/module_I_ligand_ranks.csv"),
        ("S17_Concordance_Counts", ROOT / "results/jcim_stage2_figures/module_I_concordance_counts.csv"),
    ]
    book = Workbook()
    book.remove(book.active)
    for name, path in sheets:
        fields, rows = read_csv(path)
        write_sheet(book.create_sheet(name[:31]), fields, rows)
    book.save(MAN / "Supplementary_Tables.xlsx")

    package_files = [path for _name, path in sheets]
    package_files.extend(
        [
            ROOT / "results/jcim_stage2/RUN_MANIFEST.json",
            ROOT / "reruns/UNIFORM_RERUN_V4_2_20260921/13_qa/PRIMARY_DIRECTIONAL_ANALYSIS_POPULATION.csv",
            ROOT / "reruns/UNIFORM_RERUN_V4_2_20260921/00_protocol/parent_alias_registry.csv",
            ROOT / "results/canonical/model_fold_assignments.csv",
            ROOT / "reruns/UNIFORM_RERUN_V4_2_20260921/00_protocol/FORMAL_AUTHORITY_PATHS.yaml",
            ROOT / "reruns/UNIFORM_RERUN_V4_2_20260921/00_protocol/FORMAL_METRICS_ANALYSIS_FREEZE_FINAL.yaml",
            ROOT / "reruns/UNIFORM_RERUN_V4_2_20260921/00_protocol/SCORING_DOCKING_ABLATION_FREEZE_FINAL.yaml",
            ROOT / "reruns/UNIFORM_RERUN_V4_2_20260921/00_protocol/ALTERNATIVE_RECEPTOR_EXPERIMENT_FREEZE_FINAL.yaml",
            ROOT / "reruns/UNIFORM_RERUN_V4_2_20260921/13_qa/official_primary_seed42_score_master_8pair.csv",
            ROOT / "reruns/UNIFORM_RERUN_V4_2_20260921/13_qa/phase7_fiveseed_master.csv",
            ROOT / "reruns/UNIFORM_RERUN_V4_2_20260921/13_qa/alt_score_master_independent.csv",
        ]
    )
    manifest_path = VIEWS / "SHA256SUMS.csv"
    with manifest_path.open("w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["relative_path", "sha256", "bytes"])
        for path in package_files:
            rel = path.relative_to(REPO).as_posix()
            writer.writerow([rel, sha256(path), path.stat().st_size])
    zip_path = MAN / "Supplementary_Source_Data.zip"
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in package_files:
            archive.write(path, path.relative_to(REPO).as_posix())
        archive.write(manifest_path, "SHA256SUMS.csv")
    print("TABLES_AND_PACKAGE", zip_path, "mapping", len(mapping), "population", len(pop))
    build_review_pdfs()


if __name__ == "__main__":
    main()
